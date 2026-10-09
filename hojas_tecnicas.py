"""Módulo complementario: cálculo auditable y reporte. Uso bajo revisión del responsable técnico."""
import io, math, json
from datetime import date
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER

LAB="Laboratorio de Suelos & Concreto Misoft E.I.R.L."
UNIDADES={
 'Registro de excavación': ['Profundidad inicio (m)','Profundidad fin (m)','Descripción estrato','Humedad/consistencia'],
 'Clasificación de suelo':['Tamiz 200 pasante (%)','Límite líquido (%)','Límite plástico (%)'],
 'Granulometría':['Abertura (mm)','Masa retenida (g)'],
 'Humedad':['Tara (g)','Húmedo + tara (g)','Seco + tara (g)'],
 'Límites':['Golpes (N)','Húmedo + tara (g)','Seco + tara (g)','Tara (g)'],
 'Proctor':['Contenido humedad (%)','Densidad húmeda (g/cm³)'],
 'Corte directo':['Esfuerzo normal (kPa)','Esfuerzo cortante pico (kPa)'],
 'Capacidad portante':['Profundidad desplante Df (m)','Ancho cimiento B (m)','Cohesión c (kPa)','Peso unitario γ (kN/m³)','Ángulo fricción φ (°)','Factor seguridad'],
 'Sales':['Analito','Resultado','Unidad']}
NUMERICAS={'Clasificación de suelo','Granulometría','Humedad','Límites','Proctor','Corte directo','Capacidad portante'}

def _num(x):
 try:
  f=float(x)
  return f if math.isfinite(f) else None
 except (TypeError,ValueError):return None

def calcular(tipo,df):
 """Cálculos algebraicos, no suponen acreditación ni completitud normativa."""
 resultados={}; aviso=[]; tabla=df.copy()
 if tipo=='Humedad':
  vals=[]
  for _,r in tabla.iterrows():
   t,h,s=(_num(r[c]) for c in UNIDADES[tipo]); v=None
   if None not in (t,h,s) and s>t and h>=s:v=100*(h-s)/(s-t)
   vals.append(round(v,3) if v is not None else None)
  tabla['w (%)']=vals
  v=[x for x in vals if x is not None]
  if v:resultados['Humedad promedio (%)']=round(sum(v)/len(v),3)
 elif tipo=='Granulometría':
  if len(tabla):
   v=tabla.copy();v['Abertura (mm)']=pd.to_numeric(v['Abertura (mm)'],errors='coerce');v['Masa retenida (g)']=pd.to_numeric(v['Masa retenida (g)'],errors='coerce')
   v=v.dropna().sort_values('Abertura (mm)',ascending=False)
   if len(v) and (v['Masa retenida (g)']>=0).all() and v['Masa retenida (g)'].sum()>0:
    total=v['Masa retenida (g)'].sum();v['Retenido acumulado (%)']=100*v['Masa retenida (g)'].cumsum()/total;v['Pasa (%)']=100-v['Retenido acumulado (%)'];tabla=v.round(3);resultados['Masa retenida total (g)']=round(float(total),3);aviso.append('Pasa (%) normalizado por masa retenida total; verificar masa inicial, lavado, pérdidas y fondo antes de aprobar.')
 elif tipo=='Límites':
  ws=[]
  for _,r in tabla.iterrows():
   n,h,s,t=(_num(r[c]) for c in UNIDADES[tipo]);w=100*(h-s)/(s-t) if None not in (h,s,t) and s>t and h>=s else None
   ws.append(round(w,3) if w is not None else None)
  tabla['w (%)']=ws
  pairs=[(_num(r['Golpes (N)']),_num(r['w (%)'])) for _,r in tabla.iterrows()]
  pairs=[(n,w) for n,w in pairs if n is not None and w is not None and n>0]
  if len(pairs)>=2 and len({n for n,w in pairs})>=2:
   import numpy as np
   a,b=np.polyfit([math.log10(n) for n,w in pairs],[w for n,w in pairs],1)
   resultados['LL interpolado a 25 golpes (%)']=round(a*math.log10(25)+b,2)
   aviso.append('Corresponde únicamente a aproximación semilogarítmica de límite líquido; LP debe medirse por separado.')
 elif tipo=='Clasificación de suelo':
  if len(tabla):
   r=tabla.iloc[0];ll=_num(r['Límite líquido (%)']);lp=_num(r['Límite plástico (%)'])
   if None not in (ll,lp) and ll>=lp:resultados['IP (%)']=round(ll-lp,3)
  aviso.append('No se asigna SUCS/AASHTO automáticamente sin fracciones granulométricas y verificación del método.')
 elif tipo=='Proctor':
  vs=[]
  for _,r in tabla.iterrows():
   w,g=(_num(r[c]) for c in UNIDADES[tipo]);vs.append(round(g/(1+w/100),4) if w is not None and g is not None and w>=0 and g>0 else None)
  tabla['Densidad seca (g/cm³)']=vs
  pares=[(_num(r['Contenido humedad (%)']),_num(r['Densidad seca (g/cm³)'])) for _,r in tabla.iterrows()];pares=[x for x in pares if None not in x]
  if pares:
   w,g=max(pares,key=lambda a:a[1]);resultados['Mayor densidad seca MEDIDA (g/cm³)']=g;resultados['Humedad asociada MEDIDA (%)']=w
   aviso.append('Máximo observado, NO MDS/OMC definitiva; requiere ajuste de curva, controles de compactación, energía y método.')
 elif tipo=='Corte directo':
  pts=[(_num(r['Esfuerzo normal (kPa)']),_num(r['Esfuerzo cortante pico (kPa)'])) for _,r in tabla.iterrows()];pts=[x for x in pts if None not in x and x[0]>=0 and x[1]>=0]
  if len(pts)>=2 and len(set(x for x,y in pts))>1:
   import numpy as np
   slope,intercept=np.polyfit([p[0] for p in pts],[p[1] for p in pts],1)
   resultados['c ajustada (kPa)']=round(float(intercept),3);resultados['φ ajustado (°)']=round(math.degrees(math.atan(slope)),3)
   aviso.append('Ajuste lineal de valores pico registrados: requiere desplazamientos, áreas corregidas, drenaje y validación del especialista.')
 elif tipo=='Capacidad portante':
  aviso.append('Este módulo registra parámetros: NO calcula q admisible de diseño sin geometría, nivel freático, condición drenada, excentricidad, asentamientos y método sustentado.')
 elif tipo=='Sales':aviso.append('Registrar método y unidades; la interpretación depende de límites especificados para cada analito.')
 elif tipo=='Registro de excavación':aviso.append('Describir estratigrafía, evidencia fotográfica y nivel freático en las observaciones.')
 return tabla,resultados,aviso

def grafica(tipo,tabla):
 cols={'Granulometría':('Abertura (mm)','Pasa (%)'),'Proctor':('Contenido humedad (%)','Densidad seca (g/cm³)'),'Corte directo':('Esfuerzo normal (kPa)','Esfuerzo cortante pico (kPa)'),'Límites':('Golpes (N)','w (%)')}
 if tipo not in cols:return None
 a,b=cols[tipo]
 if a not in tabla or b not in tabla:return None
 x=pd.to_numeric(tabla[a],errors='coerce');y=pd.to_numeric(tabla[b],errors='coerce');pts=[(v,w) for v,w in zip(x,y) if pd.notna(v) and pd.notna(w) and (v>0 if tipo in ['Granulometría','Límites'] else True)]
 if not pts:return None
 pts.sort();fig,ax=plt.subplots(figsize=(7.0,3.4));ax.plot([p[0] for p in pts],[p[1] for p in pts],marker='o');ax.set_xlabel(a);ax.set_ylabel(b);ax.grid(alpha=.3)
 if tipo in ['Granulometría','Límites']:ax.set_xscale('log')
 if tipo=='Granulometría':ax.set_ylim(0,100)
 fig.tight_layout();buff=io.BytesIO();fig.savefig(buff,format='png',dpi=170);plt.close(fig);buff.seek(0);return buff.getvalue()

def informe_docx(tipo,codigo,sede,tabla,resultados,avisos,tecnico,norma,obs,imagen):
 doc=Document();sec=doc.sections[0];sec.header.paragraphs[0].text=LAB+' | INFORME DE LABORATORIO'
 title=doc.add_heading('HOJA TÉCNICA DE '+tipo.upper(),0)
 for lbl,val in [('Muestra / probeta',codigo),('Sede',sede),('Fecha de emisión',date.today().isoformat()),('Técnico responsable',tecnico or 'No consignado'),('Método informado',norma or 'No consignado'),('Estado','BORRADOR – pendiente de revisión técnica')]:doc.add_paragraph(lbl+': '+str(val))
 doc.add_heading('Mediciones',level=1);t=doc.add_table(rows=1,cols=len(tabla.columns));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,tabla.columns):c.text=str(v)
 for _,r in tabla.iterrows():
  cells=t.add_row().cells
  for c,v in zip(cells,r.tolist()):c.text='' if pd.isna(v) else str(v)
 doc.add_heading('Resultados automáticos',level=1)
 if resultados:
  for k,v in resultados.items():doc.add_paragraph(f'{k}: {v}',style='List Bullet')
 else:doc.add_paragraph('Sin resultados automáticos suficientes.')
 if imagen:doc.add_picture(io.BytesIO(imagen),width=Cm(15))
 doc.add_heading('Observaciones y restricciones',level=1)
 for a in avisos:doc.add_paragraph(a,style='List Bullet')
 doc.add_paragraph('Observaciones: '+(obs or 'No consignadas'))
 doc.add_paragraph('Documento de trabajo. No constituye certificado acreditado ni informe firmado. Revisar datos, equipos, método y criterios de aceptación.')
 doc.add_paragraph('Revisó y aprobó: ________________________    Fecha: __________')
 buff=io.BytesIO();doc.save(buff);return buff.getvalue()

def informe_pdf(tipo,codigo,sede,tabla,resultados,avisos,tecnico,norma,obs,imagen):
 buff=io.BytesIO();sty=getSampleStyleSheet();normal=sty['Normal'];elements=[Paragraph(LAB,sty['Heading1']),Paragraph('HOJA TÉCNICA DE '+tipo.upper(),sty['Heading2'])]
 for k,v in [('Muestra',codigo),('Sede',sede),('Fecha',date.today()),('Técnico',tecnico or 'No consignado'),('Norma/método',norma or 'No consignado'),('Estado','BORRADOR – pendiente de revisión')]:elements.append(Paragraph(f'<b>{k}:</b> {str(v).replace("&","&amp;")}',normal))
 elements.append(Spacer(1,14));elements.append(Paragraph('Mediciones',sty['Heading2']))
 # Dividir tablas amplias en bloques legibles en página horizontal no disponible: tipografía pequeña.
 vals=[list(map(str,tabla.columns))]+[[('' if pd.isna(x) else str(x)) for x in row] for row in tabla.itertuples(index=False,name=None)]
 if len(vals)>35:vals=vals[:35];elements.append(Paragraph('Vista acotada a 34 filas; revisar fuente digital para conjunto completo.',normal))
 from reportlab.lib.pagesizes import A4,landscape
 width=landscape(A4)[0]-65
 t=Table(vals,repeatRows=1,colWidths=[width/max(1,len(tabla.columns))]*len(tabla.columns));t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#163c58')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),7),('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),5)]));elements.append(t)
 elements.append(Spacer(1,12));elements.append(Paragraph('Resultados',sty['Heading2']))
 for k,v in resultados.items():elements.append(Paragraph(f'{k}: {v}',normal))
 if imagen:elements.append(Spacer(1,8));elements.append(Image(io.BytesIO(imagen),width=410,height=200))
 elements.append(Paragraph('Observaciones y restricciones',sty['Heading2']))
 for a in avisos:elements.append(Paragraph('• '+a,normal))
 elements.append(Paragraph('Observaciones: '+(obs or 'No consignadas').replace('&','&amp;'),normal))
 elements.append(Spacer(1,12));elements.append(Paragraph('BORRADOR. Verificar resultados, procedimientos, equipos y firma del responsable antes de emisión oficial.',normal))
 elements.append(Paragraph('Revisó y aprobó: ______________________',normal))
 SimpleDocTemplate(buff,pagesize=landscape(A4),leftMargin=32,rightMargin=32,topMargin=30,bottomMargin=30).build(elements);return buff.getvalue()

def mostrar_hojas(st,client,perfil,muestras,sucursales,sucursal_filtro):
 st.header('Hojas técnicas de medición, gráficos e informes')
 st.warning('Cálculos preliminares. Los reportes se generan como BORRADORES hasta revisión del responsable del ensayo; no implican acreditación.')
 tab1,tab2=st.tabs(['Suelos','Concreto'])
 with tab1:
  if not muestras:st.info('Registre primero una muestra de suelo.');return
  tipo=st.selectbox('Ensayo / ficha',list(UNIDADES))
  opts={f"{m['codigo']} · {m.get('proyecto','')}":m for m in muestras};m=opts[st.selectbox('Muestra asociada',list(opts))]
  nombre_sede=next((s['nombre'] for s in sucursales if s['id']==m['sucursal_id']),'Sede histórica')
  try: rows=client.table('fichas_suelo').select('*').eq('muestra_id',m['id']).eq('tipo',tipo).order('creado_en',desc=True).execute().data or []
  except Exception as ex:st.error(f'No se pudo consultar fichas: {ex}');return
  options=['Nueva hoja']+[f"{r['id'][:8]} · {r['estado']}" for r in rows];sel=st.selectbox('Cargar hoja existente o crear',options)
  old=next((r for r in rows if sel.startswith(r['id'][:8])),None)
  datos=(old or {}).get('datos') or {};saved=datos.get('tabla_mediciones')
  columnas=UNIDADES[tipo]
  inicial=pd.DataFrame(saved,columns=columnas) if isinstance(saved,list) else pd.DataFrame([{c:None for c in columnas} for _ in range(4)])
  st.caption('Agregue filas para cada medición. No rellene campos sin datos reales.')
  df=st.data_editor(inicial,num_rows='dynamic',use_container_width=True,key='editor_'+tipo+'_'+sel)
  df=df.dropna(how='all').reset_index(drop=True)
  tabla,resultados,avisos=calcular(tipo,df)
  st.subheader('Tabla calculada');st.dataframe(tabla,use_container_width=True,hide_index=True)
  if resultados:st.json(resultados)
  img=grafica(tipo,tabla)
  if img:st.image(img,caption='Gráfico generado con mediciones registradas')
  for txt in avisos:st.caption('• '+txt)
  norma=st.text_input('Norma / edición aplicable',value=(old or {}).get('norma') or '',key='norma_'+tipo)
  tecnico=st.text_input('Técnico',value=(old or {}).get('tecnico') or '',key='tecnico_'+tipo)
  obs=st.text_area('Observaciones',value=(old or {}).get('observaciones') or '',key='obs_'+tipo)
  estado=st.selectbox('Estado', ['Borrador','En revisión','Validado'],index=['Borrador','En revisión','Validado'].index((old or {}).get('estado','Borrador')),key='estado_'+tipo)
  if st.button('Guardar hoja técnica',type='primary'):
   if not len(tabla):st.error('Registre al menos una medición.')
   elif not tecnico.strip():st.error('Indique el técnico responsable.')
   elif estado=='Validado' and not norma.strip():st.error('La validación requiere método/edición.')
   else:
    registros=df.where(pd.notna(df),None).to_dict('records');payload={'muestra_id':m['id'],'tipo':tipo,'datos':dict(datos,tabla_mediciones=registros,resultados_automaticos=resultados),'norma':norma,'tecnico':tecnico,'estado':estado,'observaciones':obs}
    try:
     if old:client.table('fichas_suelo').update(payload).eq('id',old['id']).execute()
     else:client.table('fichas_suelo').insert(payload).execute()
     st.success('Mediciones guardadas en ficha asociada a muestra y sede.')
    except Exception as ex:st.error(f'No se pudo guardar: {ex}')
  if len(tabla):
   xbuf=io.BytesIO()
   with pd.ExcelWriter(xbuf,engine='openpyxl') as writer:
    tabla.to_excel(writer,index=False,sheet_name='Mediciones')
    pd.DataFrame(list(resultados.items()),columns=['Resultado','Valor']).to_excel(writer,index=False,sheet_name='Resultados')
   st.download_button('Descargar hoja Excel',xbuf.getvalue(),file_name='MISOFT_'+tipo.replace(' ','_')+'_'+m['codigo']+'.xlsx')
   st.download_button('Informe Word',informe_docx(tipo,m['codigo'],nombre_sede,tabla,resultados,avisos,tecnico,norma,obs,img),file_name='MISOFT_'+m['codigo']+'.docx',mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
   st.download_button('Informe PDF',informe_pdf(tipo,m['codigo'],nombre_sede,tabla,resultados,avisos,tecnico,norma,obs,img),file_name='MISOFT_'+m['codigo']+'.pdf',mime='application/pdf')
 with tab2:
  st.subheader('Resistencia a compresión de probetas')
  st.caption('Acceso a probetas registradas en el módulo Ensayos de concreto. Seleccione una probeta para obtener cálculo verificable y reporte.')
  try: ps=client.table('probetas_concreto').select('*').order('creado_en',desc=True).limit(1000).execute().data or []
  except Exception as ex:st.error(f'Consulta de probetas no disponible: {ex}');return
  if sucursal_filtro:ps=[p for p in ps if p['sucursal_id']==sucursal_filtro]
  if not ps:st.info('Registre una probeta en Ensayos de concreto.');return
  choices={f"{p['codigo']} · {p.get('proyecto','')}":p for p in ps};p=choices[st.selectbox('Probeta',list(choices))]
  diam=_num(p.get('diametro_mm'));kn=_num(p.get('carga_kn'));geo=p.get('geometria')
  if geo=='Cilindro' and diam and diam>0 and kn and kn>0: area=math.pi*diam**2/4
  elif geo=='Cubo' and diam and diam>0 and kn and kn>0:area=diam**2
  else:area=None
  r=round(kn*1000/area,3) if area else None
  result={'Resistencia calculada (MPa)':r} if r is not None else {}
  df=pd.DataFrame([{'Geometría':geo,'Dimensión transversal (mm)':diam,'Carga rotura (kN)':kn,'Área (mm²)':round(area,2) if area else None,'f (MPa)':r}])
  st.dataframe(df,hide_index=True,use_container_width=True)
  if r is not None:
   sede=next((s['nombre'] for s in sucursales if s['id']==p['sucursal_id']),'Sede histórica')
   warning=['Se requiere registro de dimensiones efectivo, edad, curado, calibración, corrección geométrica y método aplicable.']
   st.download_button('Informe concreto Word',informe_docx('Compresión de concreto',p['codigo'],sede,df,result,warning,p.get('tecnico'),p.get('norma'),p.get('observaciones'),None),file_name=p['codigo']+'_concreto.docx')
   st.download_button('Informe concreto PDF',informe_pdf('Compresión de concreto',p['codigo'],sede,df,result,warning,p.get('tecnico'),p.get('norma'),p.get('observaciones'),None),file_name=p['codigo']+'_concreto.pdf')
